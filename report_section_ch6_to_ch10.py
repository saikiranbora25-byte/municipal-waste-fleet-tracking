# -*- coding: utf-8 -*-
"""Chapters 6 to 10: DFDs, System Architecture, Modules, Use Cases, UML Design"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_chapters_6_to_10(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_paragraph = helpers['add_paragraph']
    add_bullet = helpers['add_bullet']
    add_figure = helpers['add_figure']
    add_table_custom = helpers['add_table_custom']

    # ==========================================
    # CHAPTER 6: DATA FLOW DIAGRAMS
    # ==========================================
    add_heading_1("6. Data Flow Diagram")
    add_paragraph(
        "A Data Flow Diagram (DFD) provides a structured graphical representation of information flow through the CivicClean "
        "system, illustrating how input data from external entities is transformed through discrete logical processes into outputs "
        "and persistent data stores. DFDs are organized hierarchically into Level 0 (Context Level), Level 1, and Level 2 diagrams."
    )

    add_heading_2("6.1 Level 0 Context DFD")
    add_paragraph(
        "The Level 0 Context Diagram establishes the highest-level architectural boundary of the CivicClean system. It models the "
        "entire application as a single centralized software process (0.0) interacting with three primary external entities: the "
        "Collection Truck Driver (Field Staff), the Municipal Sanitation Supervisor, and the City Resident / Citizen."
    )
    add_figure(
        'report_assets/dfd_level_0.png',
        "Figure 6.1: Level 0 Context Data Flow Diagram (DFD)",
        "Short Explanation: Figure 6.1 illustrates the context-level data flows of the CivicClean system. The Collection Truck Driver "
        "transmits real-time GPS telemetry and geotagged collection proof photos while receiving daily assigned route manifests. The "
        "Sanitation Supervisor receives live verification audit logs, telemetry streams, and clearance analytics, while issuing route "
        "schedules and exception review decisions. The Citizen entity submits geotagged waste overflow reports and receives tracking "
        "ticket IDs and live pickup status updates."
    )

    add_heading_2("6.2 Level 1 DFD")
    add_paragraph(
        "The Level 1 Data Flow Diagram decomposes the system into five core functional sub-processes and four persistent data stores. "
        "It illustrates the internal routing of coordinates, proof imagery, grievance tickets, and audit trails."
    )
    add_figure(
        'report_assets/dfd_level_1.png',
        "Figure 6.2: Level 1 Data Flow Diagram (DFD)",
        "Short Explanation: Figure 6.2 models the five primary processes of the system: (1.0) Fleet Tracking & Telemetry Engine, "
        "(2.0) Checkpoint Geofencing & Route Engine, (3.0) Proof-of-Service Verification Engine, (4.0) Citizen Grievance Redressal "
        "Module, and (5.0) Audit & Reporting Module. Data persists across four dedicated repositories: D1 (Vehicles & Live Telemetry), "
        "D2 (Checkpoints & Route Master), D3 (Proof Logs & Audit Trail), and D4 (Citizen Grievances DB)."
    )
    add_paragraph(
        "In this flow, driver telemetry feeds into Process 1.0 to update vehicle coordinates in D1. Process 2.0 cross-references "
        "vehicle positions with checkpoint geometries in D2. Process 3.0 captures photo proof and timestamps, persisting verified "
        "records into D3 for supervisor review. Citizen reports flow through Process 4.0 into D4, dynamically linking to route dispatches."
    )

    add_heading_2("6.3 Level 2 DFD (Decomposition of Process 3.0)")
    add_paragraph(
        "To provide granular insight into the core proof-of-service verification logic, Process 3.0 is decomposed into its sub-processes "
        "in the Level 2 DFD, shown in Figure 6.3."
    )
    add_figure(
        'report_assets/dfd_level_2.png',
        "Figure 6.3: Level 2 DFD (Process 3.0 Proof-of-Service Decomposition)",
        "Short Explanation: Figure 6.3 details the internal logic of the Proof-of-Service Verification Engine. Sub-process 3.1 performs "
        "a Haversine proximity check between the truck's device GPS and the pre-registered bin coordinates. If within 15 meters, Sub-process "
        "3.2 validates the camera photo hash and device timestamp. Sub-process 3.3 captures exception justifications if a stop is skipped. "
        "Sub-process 3.4 records waste weight estimates, updates the checkpoint status to 'Collected', and writes immutable entries to the "
        "supervisor audit log."
    )

    doc.add_page_break()

    # ==========================================
    # CHAPTER 7: SYSTEM ARCHITECTURE
    # ==========================================
    add_heading_1("7. System Architecture")
    
    add_heading_2("7.1 3-Tier Client-Server Architecture")
    add_paragraph(
        "The CivicClean system is built on a clean 3-Tier Client-Server Architecture comprising the Presentation Tier, the Application "
        "(Business Logic) Tier, and the Data Persistence Tier. This architectural separation enforces high modularity, scalability, and "
        "maintainability, while enabling seamless software-only deployment."
    )
    add_figure(
        'report_assets/system_architecture.png',
        "Figure 7.1: 3-Tier System Architecture Diagram",
        "Short Explanation: Figure 7.1 illustrates the structural tiers of the CivicClean platform. The Presentation Tier accommodates "
        "desktop and mobile web clients for supervisors, drivers, and citizens. The Application Tier executes on a Node.js/Express runtime, "
        "housing the REST API Gateway, GIS Geofencing Engine, Proof Verification Service, and Grievance Controller. The Data Tier maintains "
        "vehicle rosters, route geometries, immutable audit logs, and citizen grievance media."
    )

    add_heading_2("7.2 Architectural Layer Analysis")
    add_bullet("Presentation Layer (Client Interfaces)", "Constructed using responsive HTML5, modern CSS3 with mobile shell framing, and vanilla JavaScript. It delivers role-tailored views: the Supervisor GIS Map powered by Leaflet.js, the touch-optimized Driver Field PWA, and the public Citizen Grievance Portal. No proprietary client software installation is required; any standard browser suffices.")
    add_bullet("Application Layer (Business Logic & REST APIs)", "Hosted on an asynchronous Node.js engine with Express middleware. This layer encapsulates the core business rules: spatial proximity computation using the spherical Haversine formula, tamper-evident timestamp validation, vehicle GPS simulation for testing, and automated ticket dispatch algorithms. Communication with clients is handled via stateless HTTP/JSON endpoints.")
    add_bullet("Data Persistence Layer", "Provides structured storage for operational states, including vehicle manifests, scheduled stop sequences, photo proof records, and audit logs. The in-memory data store with disk persistence ensures rapid sub-millisecond retrieval during live tracking, while preserving complete audit logs for administrative compliance.")

    doc.add_page_break()

    # ==========================================
    # CHAPTER 8: MODULE DESCRIPTION
    # ==========================================
    add_heading_1("8. Module Description")
    add_paragraph(
        "The CivicClean system is partitioned into four tightly cohesive, loosely coupled functional modules, each addressing "
        "the requirements of specific municipal stakeholders."
    )

    add_heading_2("8.1 Supervisor Command Hub & GIS Fleet Tracking Module")
    add_paragraph(
        "This module serves as the primary operational console for municipal sanitation superintendents and route inspectors. Key capabilities include:"
    )
    add_bullet("Interactive Leaflet GIS Canvas", "Renders OpenStreetMap vector tiles covering the municipal ward sector (e.g., Sangivalasa to Bheemili). Route polylines are dynamically drawn, and collection checkpoints are displayed with distinct visual markers (green for collected, amber for pending, red for skipped, and purple for citizen complaints).")
    add_bullet("Real-Time Vehicle Telemetry", "Monitors active collection vehicles, displaying live speed (km/h), fuel level, battery health, current assigned driver, and GPS lock precision (?4m).")
    add_bullet("Simulation & Playback Controller", "Includes playback controls (Play, Pause, Step Next, Reset) allowing supervisors to simulate GPS coordinate feeds during driver training or system verification.")

    add_heading_2("8.2 Driver Mobile Field Manifest & Geofencing Module")
    add_paragraph(
        "Designed specifically for field sanitation workers operating compactor trucks, this module provides an intuitive, touch-friendly "
        "manifest interface accessible on commodity Android smartphones:"
    )
    add_bullet("Chronological Checkpoint Checklist", "Displays daily assigned stops in optimal collection order with target arrival times, ward names, and completion status badges.")
    add_bullet("GPS Lock & Geofencing Banner", "Monitors the smartphone's built-in GNSS chip, verifying continuous signal lock and calculating real-time distance to the next scheduled checkpoint.")
    add_bullet("Shift Progress Tracking", "Displays a dynamic visual completion bar (e.g., '2 / 6 Stops (33%)') updating instantly as stops are cleared.")

    add_heading_2("8.3 Proof-of-Service Verification & Tamper-Resistant Audit Module")
    add_paragraph(
        "This module enforces accountability and eliminates falsification of waste collection claims:"
    )
    add_bullet("Geotagged Camera Capture", "Upon arriving within 15 meters of a bin, the driver taps 'Mark Collected'. The app captures a photo of the cleared bin, stamping the image with the exact latitude, longitude, and server-validated timestamp.")
    add_bullet("Skip-Exception Recording", "If physical barriers (road excavation, illegal parking, hazardous waste) prevent collection, the driver opens a structured modal, selects the obstacle category, and logs observations. The system immediately flags the stop for supervisor intervention.")
    add_bullet("Immutable Audit Trail Table", "Maintains an unalterable, chronological log of all collection events, logging vehicle registration numbers, stop names, action types (COLLECTION_VERIFIED or STOP_SKIPPED), and GPS accuracy tolerances.")

    add_heading_2("8.4 Citizen Grievance Redressal & Hotspot Triage Module")
    add_paragraph(
        "This public-facing module empowers urban residents to participate actively in neighborhood cleanliness:"
    )
    add_bullet("Geo-Tagged Complaint Submission", "Residents report overflowing public bins, missed pickups, or illegal roadside dumping by providing contact details, selecting their municipal ward, uploading a photo, and pinning their exact location via GPS.")
    add_bullet("Real-Time Grievance Tracker", "Generates a unique tracking ticket (e.g., CMP-2026-104) and presents a 3-stage visual timeline showing progress from submission to truck dispatch and final photographic clearance.")
    add_bullet("Ward Collection Roster", "Displays scheduled daily pickup timings and assigned collection vehicles across wards to keep citizens informed.")

    doc.add_page_break()

    # ==========================================
    # CHAPTER 9: SYSTEM ANALYSIS
    # ==========================================
    add_heading_1("9. System Analysis")
    
    add_heading_2("9.1 Use Case Diagram")
    add_paragraph(
        "The Use Case Diagram models the functional interactions between external human actors and the CivicClean system boundary. "
        "The primary actors are the Driver (Field Staff), the Citizen (Resident), and the Sanitation Supervisor."
    )
    add_figure(
        'report_assets/use_case_diagram.png',
        "Figure 9.1: System Use Case Diagram",
        "Short Explanation: Figure 9.1 illustrates the system use case interactions. The Driver interacts with UC1 (View Route Manifest), "
        "UC2 (Transmit GPS Telemetry), UC3 (Mark Stop Collected with Photo), and UC4 (Log Skipped Stop with Reason). The Citizen interacts "
        "with UC5 (Report Waste Hotspot) and UC6 (Track Grievance Status). The Supervisor interacts with UC7 (Verify Proof-of-Service Audit) "
        "and UC8 (Monitor Live Fleet GIS Map)."
    )

    add_heading_2("9.2 Use Case Specifications")
    add_paragraph(
        "To provide rigorous engineering documentation, formal use case specifications for the three most critical system transactions "
        "are detailed in Tables 9.1, 9.2, and 9.3."
    )

    # Table 9.1: UC1
    uc1_headers = ["Use Case Element", "Specification Details"]
    uc1_rows = [
        ["Use Case ID & Name", "UC1: View Assigned Route Manifest"],
        ["Primary Actor", "Collection Truck Driver (Field Staff)"],
        ["Preconditions", "Driver has successfully authenticated into the mobile field portal; active shift assigned."],
        ["Trigger", "Driver selects the 'Driver Field App' tab upon starting shift."],
        ["Main Success Scenario", "1. Driver opens field app.\n2. System queries /api/stops for driver's assigned route ID.\n3. System renders chronological stop manifest with target times and status badges.\n4. Driver reviews checkpoints and commences travel."],
        ["Alternative / Exception Flow", "If server is unreachable, app renders cached local route manifest and alerts driver of offline mode."],
        ["Postconditions", "Driver displays verified list of today's collection stops with active GPS geofence tracking."]
    ]
    add_table_custom("Table 9.1: Use Case Specification: UC1 View Assigned Route Manifest", uc1_headers, uc1_rows, [2.2, 4.6])

    # Table 9.2: UC3
    uc2_headers = ["Use Case Element", "Specification Details"]
    uc2_rows = [
        ["Use Case ID & Name", "UC3: Mark Checkpoint Collected with Photo Proof"],
        ["Primary Actor", "Collection Truck Driver (Field Staff)"],
        ["Preconditions", "Vehicle is physically positioned within the checkpoint's 15m geofence radius; stop status is 'Pending'."],
        ["Trigger", "Driver taps 'Mark Collected' button on the active checkpoint card."],
        ["Main Success Scenario", "1. Driver taps 'Mark Collected'.\n2. App triggers camera dialog and captures bin photo.\n3. App acquires current device coordinates (lat, lng) and system timestamp.\n4. Client dispatches POST request to /api/stops/:id/collect.\n5. Server validates geofence proximity (Haversine distance <= 15m).\n6. Server updates stop status to 'Collected', records estimated weight, and appends audit log.\n7. Mobile UI updates card with green checkmark and advances progress bar."],
        ["Alternative / Exception Flow", "If truck is outside 15m geofence, system blocks completion and displays: 'Error: Must be within 15m of bin location'."],
        ["Postconditions", "Stop is verified 'Collected', proof photo is archived, and supervisor dashboard updates in real time."]
    ]
    add_table_custom("Table 9.2: Use Case Specification: UC3 Mark Stop Collected with Proof", uc2_headers, uc2_rows, [2.2, 4.6])

    # Table 9.3: UC5
    uc3_headers = ["Use Case Element", "Specification Details"]
    uc3_rows = [
        ["Use Case ID & Name", "UC5: Report Waste Hotspot or Missed Pickup"],
        ["Primary Actor", "City Resident / Citizen"],
        ["Preconditions", "Citizen has access to mobile or desktop web browser; internet connectivity active."],
        ["Trigger", "Citizen identifies an overflowing public bin or missed collection and opens the Citizen Portal."],
        ["Main Success Scenario", "1. Citizen fills name, phone number, ward, and issue category.\n2. Citizen taps 'GPS Pin' to acquire exact device location or types landmark.\n3. Citizen attaches photo of overflowing waste and submits form.\n4. Client posts payload to /api/complaints.\n5. Server stores complaint in D4 and generates tracking ticket (e.g., #CMP-2026-104).\n6. Portal displays confirmation modal with ticket number and initial 'Assigned' status."],
        ["Alternative / Exception Flow", "If mandatory fields are missing, form highlights invalid inputs with validation warnings."],
        ["Postconditions", "Grievance is visible on supervisor GIS triage map and citizen can track status via ticket ID."]
    ]
    add_table_custom("Table 9.3: Use Case Specification: UC5 Report Waste Hotspot", uc3_headers, uc3_rows, [2.2, 4.6])

    doc.add_page_break()

    # ==========================================
    # CHAPTER 10: UML DESIGN
    # ==========================================
    add_heading_1("10. UML Design")
    add_paragraph(
        "Unified Modeling Language (UML) provides standard visual modeling mechanisms to capture the structural and behavioral facets "
        "of the CivicClean software architecture. This chapter presents six comprehensive UML diagrams adhering strictly to UML 2.5 standards."
    )

    add_heading_2("10.1 Class Diagram")
    add_paragraph(
        "The UML Class Diagram models the static structural design of the system, defining object-oriented classes, attributes, "
        "operations, visibility modifiers, and association relationships."
    )
    add_figure(
        'report_assets/class_diagram.png',
        "Figure 10.1: UML Class Diagram",
        "Short Explanation: Figure 10.1 depicts the structural classes of CivicClean: Vehicle, StopCheckpoint, ProofOfService, DriverUser, "
        "CitizenComplaint, and SupervisorAdmin. Association cardinalities show that one Vehicle visits multiple StopCheckpoints (1..*), "
        "each completed StopCheckpoint produces exactly one ProofOfService record (1 to 1), and a CitizenComplaint is reviewed and assigned "
        "by a SupervisorAdmin."
    )

    add_heading_2("10.2 Sequence Diagram")
    add_paragraph(
        "The UML Sequence Diagram models dynamic behavior by illustrating the chronological exchange of messages between runtime object "
        "instances during a critical transaction: the Proof-of-Service verification flow."
    )
    add_figure(
        'report_assets/sequence_diagram.png',
        "Figure 10.2: UML Sequence Diagram (Proof-of-Service Collection Verification)",
        "Short Explanation: Figure 10.2 traces the sequential execution flow when a driver collects waste. The driver initiates the action "
        "via the Mobile Field App, which transmits GPS coordinates and photo data to the Backend REST API. The Geofence Engine validates "
        "spatial proximity (4m <= 15m threshold). The Database updates the stop status and writes an immutable audit record. The server "
        "returns an HTTP 200 verification receipt to the driver and streams real-time updates to the Supervisor Dashboard."
    )

    add_heading_2("10.3 Activity Diagram")
    add_paragraph(
        "The UML Activity Diagram models the operational control flow of the Driver's daily collection shift, including conditional "
        "branching for accessible versus obstructed checkpoints."
    )
    add_figure(
        'report_assets/activity_diagram.png',
        "Figure 10.3: UML Activity Diagram (Driver Daily Collection Workflow)",
        "Short Explanation: Figure 10.3 traces the driver's workflow from shift initiation to return to depot. Upon arriving at a checkpoint, "
        "a decision node checks accessibility: if clear, the driver snaps a photo, records waste weight, and logs proof; if blocked by roadwork "
        "or traffic, the driver logs a structured skip reason with observations. Both paths rejoin to check remaining route stops."
    )

    add_heading_2("10.4 State Diagram")
    add_paragraph(
        "The UML State Diagram captures the lifecycle state transitions of a StopCheckpoint object in response to operational events."
    )
    add_figure(
        'report_assets/state_diagram.png',
        "Figure 10.4: UML State Diagram (Stop Checkpoint Lifecycle)",
        "Short Explanation: Figure 10.4 defines the states of a collection stop: beginning in SCHEDULED (pending), transitioning to "
        "IN_PROGRESS when the truck is en route, and branching to either COLLECTED (upon geofence validation and photo proof submission) or "
        "SKIPPED (if an obstruction is flagged). Both terminal states are archived into the supervisor audit database."
    )

    add_heading_2("10.5 Component Diagram")
    add_paragraph(
        "The UML Component Diagram models the software organization, defining modular executable components, interfaces, and inter-component dependencies."
    )
    add_figure(
        'report_assets/component_diagram.png',
        "Figure 10.5: UML Component Diagram",
        "Short Explanation: Figure 10.5 shows the modular organization of CivicClean: the Web Presentation Component connects via HTTP/JSON "
        "to the API Gateway Component, which coordinates the GIS & Geofencing Component, Proof Verification Component, Grievance Redressal "
        "Component, and Data Access & Storage Component."
    )

    add_heading_2("10.6 Deployment Diagram")
    add_paragraph(
        "The UML Deployment Diagram illustrates the physical runtime environment, showing how software artifacts are deployed across "
        "hardware processing nodes and communication networks."
    )
    add_figure(
        'report_assets/deployment_diagram.png',
        "Figure 10.6: UML Deployment Diagram",
        "Short Explanation: Figure 10.6 maps software execution across physical hardware: the Driver Mobile Smartphone connects via 4G LTE "
        "cellular network to the Cloud Web/API Server, the Supervisor Workstation accesses the server via high-speed broadband, and the "
        "Cloud Server communicates with the Database Server over an isolated TCP/IP link."
    )

    doc.add_page_break()
    print("Chapters 6 to 10 generated successfully.")
