# -*- coding: utf-8 -*-
"""Slides 1 to 4: Title, Introduction, Problem Statement, Root Cause Analysis"""
import os
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_slides_1_to_4(prs, helpers):
    apply_background = helpers['apply_background']
    add_header = helpers['add_header']
    add_card = helpers['add_card']
    blank_layout = helpers['blank_layout']
    C_DARK_BG = helpers['C_DARK_BG']
    C_PRIMARY = helpers['C_PRIMARY']
    C_WHITE = helpers['C_WHITE']
    C_TEXT_SUB = helpers['C_TEXT_SUB']
    C_DANGER = helpers['C_DANGER']
    C_PRIMARY_DK = helpers['C_PRIMARY_DK']

    # ==========================================================
    # SLIDE 1: TITLE SLIDE (Dark Navy Executive Theme)
    # ==========================================================
    slide1 = prs.slides.add_slide(blank_layout)
    apply_background(slide1, C_DARK_BG)

    # Top accent line
    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_PRIMARY
    top_bar.line.fill.background()

    # ANITS College Logo
    logo_path = 'report_assets/anits_logo.jpeg'
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(0.9), Inches(0.52), width=Inches(1.25))

    # College Header Text
    col_box = slide1.shapes.add_textbox(Inches(2.35), Inches(0.5), Inches(10.1), Inches(1.1))
    tf_col = col_box.text_frame
    tf_col.word_wrap = True
    
    p_c1 = tf_col.paragraphs[0]
    p_c1.text = "ANIL NEERUKONDA INSTITUTE OF TECHNOLOGY AND SCIENCES (UGC AUTONOMOUS)"
    p_c1.font.name = "Segoe UI"
    p_c1.font.size = Pt(13.5)
    p_c1.font.bold = True
    p_c1.font.color.rgb = RGBColor(226, 232, 240)

    p_c2 = tf_col.add_paragraph()
    p_c2.text = "DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING | 23CS4219 - SOFTWARE ENGINEERING LAB"
    p_c2.font.name = "Segoe UI"
    p_c2.font.size = Pt(10)
    p_c2.font.bold = True
    p_c2.font.color.rgb = RGBColor(147, 197, 253)

    # Badge
    p_badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.82), Inches(5.2), Inches(0.35))
    p_badge.fill.solid()
    p_badge.fill.fore_color.rgb = RGBColor(30, 58, 138)
    p_badge.line.color.rgb = RGBColor(59, 130, 246)
    p_badge.line.width = Pt(1)
    tf_pb = p_badge.text_frame
    tf_pb.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_pb = tf_pb.paragraphs[0]
    p_pb.text = "MINI PROJECT APPLICATION & TECHNICAL EVALUATION"
    p_pb.font.name = "Segoe UI"
    p_pb.font.size = Pt(9.5)
    p_pb.font.bold = True
    p_pb.font.color.rgb = RGBColor(191, 219, 254)

    t_box = slide1.shapes.add_textbox(Inches(0.9), Inches(2.26), Inches(11.5), Inches(1.85))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    
    p_pt = tf_t.paragraphs[0]
    p_pt.text = "GPS-Based Fleet Tracking & Route Verification System\nfor Municipal Solid Waste Management"
    p_pt.font.name = "Segoe UI"
    p_pt.font.size = Pt(25)
    p_pt.font.bold = True
    p_pt.font.color.rgb = C_WHITE
    p_pt.space_after = Pt(8)

    p_sub = tf_t.add_paragraph()
    p_sub.text = "A Hardware-Free, Software-Only Smart Waste Logistics & Proof-of-Service Verification Platform"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = RGBColor(148, 163, 184)

    # Card 1: Student Details
    c1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(4.35), Inches(5.6), Inches(2.65))
    c1.fill.solid()
    c1.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c1.line.color.rgb = RGBColor(51, 65, 85)
    c1.line.width = Pt(1.5)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = Inches(0.3)
    tf_c1.margin_top = Inches(0.2)
    
    p_h1 = tf_c1.paragraphs[0]
    p_h1.text = "PROJECT DEVELOPER & PRESENTER"
    p_h1.font.name = "Segoe UI"
    p_h1.font.size = Pt(11)
    p_h1.font.bold = True
    p_h1.font.color.rgb = RGBColor(56, 189, 248)
    p_h1.space_after = Pt(8)

    details = [
        ("Candidate Name:", "SAIKIRAN BORA"),
        ("Roll Number:", "A24126510006"),
        ("Branch / Degree:", "Computer Science & Engineering (B.Tech)"),
        ("Class & Section:", "III Year, I Semester, Section - A"),
        ("Institution:", "ANITS (Anil Neerukonda Inst of Tech & Sci)")
    ]
    for lbl, val in details:
        p_d = tf_c1.add_paragraph()
        p_d.text = f"{lbl}  {val}"
        p_d.font.name = "Segoe UI"
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = RGBColor(241, 245, 249)
        p_d.space_after = Pt(3)

    # Card 2: Guidance Details
    c2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.35), Inches(5.6), Inches(2.65))
    c2.fill.solid()
    c2.fill.fore_color.rgb = RGBColor(15, 23, 42)
    c2.line.color.rgb = RGBColor(51, 65, 85)
    c2.line.width = Pt(1.5)
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = Inches(0.3)
    tf_c2.margin_top = Inches(0.2)
    
    p_h2 = tf_c2.paragraphs[0]
    p_h2.text = "ACADEMIC SUPERVISION & GUIDANCE"
    p_h2.font.name = "Segoe UI"
    p_h2.font.size = Pt(11)
    p_h2.font.bold = True
    p_h2.font.color.rgb = RGBColor(52, 211, 153)
    p_h2.space_after = Pt(8)

    g_details = [
        ("Faculty Guide:", "Prof. A. Rohini"),
        ("Designation:", "Faculty Incharge, Dept. of CSE"),
        ("Head of Department:", "Prof. G. Srinivas"),
        ("Designation:", "HOD & Professor, Dept. of CSE, ANITS"),
        ("Academic Year:", "2026 - 2027")
    ]
    for lbl, val in g_details:
        p_gd = tf_c2.add_paragraph()
        p_gd.text = f"{lbl}  {val}"
        p_gd.font.name = "Segoe UI"
        p_gd.font.size = Pt(12)
        p_gd.font.color.rgb = RGBColor(241, 245, 249)
        p_gd.space_after = Pt(3)

    # ==========================================================
    # SLIDE 2: INTRODUCTION & CONTEXT (12pt Font)
    # ==========================================================
    slide2 = prs.slides.add_slide(blank_layout)
    apply_background(slide2)
    add_header(slide2, "Introduction: Urban Solid Waste Management Realities", "01 | DOMAIN CONTEXT", 2)

    cards_data = [
        ("1. Urban Waste Generation Scale",
         "\u2022 Rapid urban growth in cities generates over 500 metric tons of garbage every day.\n"
         "\u2022 Spans dense residential colonies, busy commercial zones, and narrow alleys.\n"
         "\u2022 Requires synchronized dispatch of municipal collection trucks every single morning.",
         C_PRIMARY),
        ("2. Public Health & Civic Hygiene",
         "\u2022 Overflowing secondary garbage bins emit severe odors and create public hazards.\n"
         "\u2022 Accelerates vector-borne disease outbreaks like dengue, malaria, and cholera.\n"
         "\u2022 Unchecked toxic leachate fluid contaminates storm drains and underground water.",
         RGBColor(16, 185, 129)),
        ("3. High Operational Expenditure",
         "\u2022 Municipal corporations spend 35% to 40% of their total annual budget on waste logistics.\n"
         "\u2022 Fuel consumption and truck fleet maintenance represent heavy daily expenses.\n"
         "\u2022 Manual pen-and-paper tracking leads to huge financial waste and missed pickups.",
         RGBColor(139, 92, 246))
    ]

    for idx, (title, text, col) in enumerate(cards_data):
        cx = Inches(0.8 + idx * 3.95)
        card, tf = add_card(slide2, cx, Inches(1.68), Inches(3.8), Inches(5.0), title, title_color=col)
        for line in text.split('\n'):
            p = tf.add_paragraph()
            p.text = line.encode('utf-8').decode('unicode_escape')
            p.font.name = "Segoe UI"
            p.font.size = Pt(12)
            p.font.color.rgb = C_TEXT_SUB
            p.space_after = Pt(6)

    # ==========================================================
    # SLIDE 3: PROBLEM STATEMENT (12pt Font)
    # ==========================================================
    slide3 = prs.slides.add_slide(blank_layout)
    apply_background(slide3)
    add_header(slide3, "Problem Statement: Breakdown of Conventional Manual Logistics", "02 | PROBLEM STATEMENT", 3)

    p_cards = [
        ("Static, Inflexible Route Sheets",
         "Collection trucks follow rigid, paper-based schedules. Vehicles are dispatched blindly without knowing which bins are actually overflowing or already serviced.",
         RGBColor(245, 158, 11)),
        ("Zero Route Verifiability & Accountability",
         "Sanitation supervisors have NO digital proof or timestamped logs to verify whether a driver actually emptied a scheduled bin or completely bypassed it.",
         C_DANGER),
        ("Substantial Fuel & Fleet Wastage",
         "Compactor trucks burn expensive diesel visiting empty bins, while heavily overflowing market checkpoints remain untouched, causing major resource waste.",
         C_DANGER),
        ("Reactive & Informal Grievance Loop",
         "Citizens discover missed pickups only after days of overflow. Complaints are handled via informal phone calls that lack ticket numbers or resolution proof.",
         RGBColor(245, 158, 11))
    ]

    for idx, (title, desc, col) in enumerate(p_cards):
        row = idx // 2
        col_idx = idx % 2
        cx = Inches(0.8 + col_idx * 5.95)
        cy = Inches(1.68 + row * 2.55)
        card, tf = add_card(slide3, cx, cy, Inches(5.75), Inches(2.38), title, border_color=col, title_color=col)
        p = tf.add_paragraph()
        p.text = desc
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_SUB
        p.space_after = Pt(4)

    # ==========================================================
    # SLIDE 4: ROOT CAUSE ANALYSIS (IoT Fallacy vs Reality - 12pt Font)
    # ==========================================================
    slide4 = prs.slides.add_slide(blank_layout)
    apply_background(slide4)
    add_header(slide4, "Root Cause Analysis: Why Previous IoT Hardware Approaches Failed", "03 | ROOT CAUSE ANALYSIS", 4)

    # Left: IoT Sensor Fallacy
    c_left, tf_l = add_card(slide4, Inches(0.8), Inches(1.68), Inches(5.75), Inches(5.0),
                            "The Fallacy of Dedicated IoT Hardware (Solution 1)",
                            border_color=C_DANGER, title_color=C_DANGER)
    iot_points = [
        ("Prohibitive Capital Cost:", "Fitting thousands of physical sensors across roadside bins costs millions in municipal funds."),
        ("Rampant Public Vandalism:", "Outdoor sensors are stolen, damaged by scavengers, or smashed by hydraulic truck compactor arms."),
        ("Harsh Environmental Damage:", "Acidic garbage leachate, heavy monsoon rain, and toxic gases rapidly destroy electronic circuitry."),
        ("Battery Replacement Burden:", "Swapping batteries across thousands of scattered city bins creates an unmanageable maintenance load.")
    ]
    for lbl, desc in iot_points:
        p = tf_l.add_paragraph()
        p.text = f"\u2022 {lbl} {desc}".encode('utf-8').decode('unicode_escape')
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_SUB
        p.space_after = Pt(8)

    # Right: The Real Root Cause
    c_right, tf_r = add_card(slide4, Inches(6.8), Inches(1.68), Inches(5.75), Inches(5.0),
                             "The Core Root Cause Identified",
                             border_color=C_PRIMARY, title_color=C_PRIMARY_DK)
    p_core = tf_r.add_paragraph()
    p_core.text = (
        "Municipal solid waste management fails NOT because roadside bins lack chips, "
        "but because there is NO real-time digital feedback loop connecting trucks, "
        "drivers, supervisors, and citizens."
    )
    p_core.font.name = "Segoe UI"
    p_core.font.size = Pt(12.5)
    p_core.font.bold = True
    p_core.font.color.rgb = C_PRIMARY_DK
    p_core.space_after = Pt(10)

    rc_points = [
        ("Missing Proof of Service:", "Paper log sheets allow post-shift falsification with zero physical verifiability."),
        ("No Obstruction Logging:", "When roads are blocked by construction, drivers bypass bins with no recorded reason."),
        ("The Strategic Takeaway:", "A pure software-only system delivers 100% verifiability and tracking without spending a single rupee on in-bin hardware!")
    ]
    for lbl, desc in rc_points:
        p = tf_r.add_paragraph()
        p.text = f"\u2022 {lbl} {desc}".encode('utf-8').decode('unicode_escape')
        p.font.name = "Segoe UI"
        p.font.size = Pt(12)
        p.font.color.rgb = C_TEXT_SUB
        p.space_after = Pt(8)

    print("Slides 1 to 4 generated successfully with 12pt fonts.")
