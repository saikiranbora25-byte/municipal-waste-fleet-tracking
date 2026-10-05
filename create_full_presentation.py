# -*- coding: utf-8 -*-
"""
CivicClean GIS Presentation Builder
Generates a 16-Slide Professional Academic PowerPoint Presentation
Author: Saikiran Bora (Roll No: A24126510006)
Branch: Computer Science & Engineering | College: ANITS
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Color Palette Definitions
C_DARK_BG   = RGBColor(11, 19, 43)      # #0B132B (Deep Navy)
C_LIGHT_BG  = RGBColor(248, 250, 252)   # #F8FAFC (Slate 50)
C_WHITE     = RGBColor(255, 255, 255)
C_PRIMARY   = RGBColor(37, 99, 235)     # #2563EB (Royal Blue)
C_PRIMARY_DK= RGBColor(30, 58, 138)     # #1E3A8A (Navy Blue)
C_SUCCESS   = RGBColor(16, 185, 129)    # #10B981 (Emerald Green)
C_DANGER    = RGBColor(239, 68, 68)     # #EF4444 (Crimson Red)
C_WARNING   = RGBColor(245, 158, 11)    # #F59E0B (Amber)
C_PURPLE    = RGBColor(139, 92, 246)    # #8B5CF6 (Purple)
C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A (Slate 900)
C_TEXT_SUB  = RGBColor(71, 85, 105)     # #475569 (Slate 600)
C_TEXT_MUTED= RGBColor(100, 116, 139)   # #64748B (Slate 500)
C_BORDER    = RGBColor(226, 232, 240)   # #E2E8F0
C_CARD_FILL = RGBColor(255, 255, 255)

def apply_background(slide, color=C_LIGHT_BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="MUNICIPAL SOLID WASTE MANAGEMENT | SOLUTION 2", slide_num=None):
    # Category Pill
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(4.5), Inches(0.32))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
    cat_box.line.color.rgb = RGBColor(191, 219, 254)
    cat_box.line.width = Pt(1)
    tf_c = cat_box.text_frame
    tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.name = "Segoe UI"
    p_c.font.size = Pt(8.5)
    p_c.font.bold = True
    p_c.font.color.rgb = C_PRIMARY
    
    # Title Text
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(21)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_MAIN

    # Subtle divider
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = C_BORDER
    div.line.fill.background()

    # Footer
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.font.name = "Segoe UI"
    p_f.font.size = Pt(8)
    p_f.font.color.rgb = C_TEXT_MUTED
    p_f.text = "Saikiran Bora (Roll No: A24126510006) | Branch: CSE, ANITS  ?  CivicClean GIS: GPS Fleet Tracking & Route Verification Platform"
    if slide_num:
        p_f.text += f"  ?  Slide {slide_num} of 16"

def add_card(slide, left, top, width, height, title="", border_color=C_BORDER, fill_color=C_CARD_FILL, title_color=C_PRIMARY_DK):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.space_after = Pt(6)
    return card, tf

print("Base presentation framework initialized.")
