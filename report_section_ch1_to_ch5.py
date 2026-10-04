# -*- coding: utf-8 -*-
"""Chapters 1 to 5: Introduction, Existing System, Proposed System, Requirements, HW/SW"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_chapters_1_to_5(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_paragraph = helpers['add_paragraph']
    add_bullet = helpers['add_bullet']
    add_table_custom = helpers['add_table_custom']

    # ==========================================
    # CHAPTER 1: INTRODUCTION
    # ==========================================
    add_heading_1("1. Introduction")
    
    add_heading_2("1.1 Background")
    add_paragraph(
        "Rapid urbanization, industrial expansion, and population growth across metropolitan and tier-2 Indian cities?such "
        "as Greater Visakhapatnam?have generated unprecedented quantities of municipal solid waste (MSW). Managing the daily "
        "collection, transportation, and safe disposal of solid waste represents one of the most resource-intensive and logistically "
        "challenging operations undertaken by municipal corporations. In a typical urban sector, hundreds of metric tons of domestic, "
        "commercial, and institutional waste must be cleared daily across dense residential wards, market yards, commercial districts, "
        "and peripheral coastal settlements. Efficient waste logistics directly determine civic public health, environmental hygiene, "
        "air and water quality, and resident satisfaction."
    )
    add_paragraph(
        "Historically, municipal solid waste management has relied almost exclusively on static, historically inherited route sheets "
        "and manual supervisory oversight. In this conventional paradigm, collection trucks are dispatched according to fixed calendar "
        "schedules without real-time knowledge of bin fill levels or driver route completion. When vehicles break down, face roadblocks, "
        "or skip designated neighborhood bins, supervisors have no mechanism to detect these anomalies until residents register formal "
        "complaints. Consequently, urban centers frequently suffer from overflowing secondary collection bins, roadside littering, "
        "severe traffic congestion caused by unoptimized truck schedules, and substantial diesel fuel waste resulting from trucks visiting "
        "low-volume sectors. Developing an automated, verifiable, and transparent software solution is therefore an imperative need "
        "for modern smart city governance."
    )

    add_heading_2("1.2 Problem Statement")
    add_paragraph(
        "The prevailing municipal waste collection tracking mechanisms suffer from critical operational deficiencies, leading to "
        "widespread service failures. Specifically, the problem encompasses the following core challenges:"
    )
    add_bullet("Absence of Route Verifiability", "Truck drivers follow informal or paper-based route manifests. Supervisors have no verifiable digital proof (such as GPS trail logs or geotagged timestamps) to ascertain whether a driver actually visited a scheduled checkpoint or skipped it.")
    add_bullet("Reactive Complaint Handling", "Municipal authorities operate reactively, discovering missed pickups only after days of resident agitation via informal phone calls that lack tracking numbers or escalation workflows.")
    add_bullet("Inefficient Fleet Resource Allocation", "Static routes cause collection vehicles to travel fixed kilometers regardless of actual waste volume, leading to excessive fuel expenditure, vehicle wear-and-tear, and avoidable carbon emissions.")
    add_bullet("Lack of Structured Exception Logging", "When drivers encounter valid physical obstructions?such as road excavations, water-logging, or inaccessible narrow lanes?there is no mobile tool to record the justification, leaving supervisors unable to distinguish between genuine operational impediments and driver negligence.")

    add_heading_2("1.3 Motivation")
    add_paragraph(
        "In recent academic literature and smart-city pilot projects, various Internet-of-Things (IoT) architectures have been "
        "proposed that install ultrasonic level sensors and wireless transmission modules (such as LoRaWAN or NB-IoT) inside public bins. "
        "However, rigorous field analysis reveals that hardware-intensive sensor deployments suffer from severe practical constraints "
        "in Indian municipal environments:"
    )
    add_bullet("Prohibitive Capital & Maintenance Expenditure", "Procuring, waterproofing, and maintaining thousands of battery-powered electronic sensors across harsh outdoor bin environments is financially unsustainable for municipal bodies operating under tight budgets.")
    add_bullet("Vandalism and Theft Vulnerability", "Public waste bins located on thoroughfares and market junctions are heavily prone to physical vandalism, accidental compaction damage by heavy garbage lifting arms, and sensor theft.")
    add_bullet("Battery Exhaustion and Harsh Environmental Wear", "Exposure to extreme heat, monsoon moisture, corrosive leachate fluids, and toxic fumes causes rapid sensor degradation and erratic telemetry.")
    add_paragraph(
        "These realities motivated the formulation of Solution 2: a pure software-only, zero-hardware deployment model. By leveraging "
        "commodity smartphones already possessed by drivers and citizens, along with modern web GIS mapping and browser APIs, the "
        "entire waste tracking and verification lifecycle can be automated at a fraction of the cost, achieving 100% service accountability "
        "without installing a single physical sensor in public bins."
    )

    add_heading_2("1.4 Objectives")
    add_paragraph(
        "The primary objective of this project is to design, develop, and evaluate a comprehensive web-based GPS Fleet Tracking "
        "and Route Verification System (CivicClean GIS) for municipal solid waste management. Specific technical objectives include:"
    )
    add_bullet("Automate Fleet Tracking", "Provide an interactive, real-time GIS map interface for sanitation supervisors to track vehicle coordinates, operational speed, and route trajectories using Leaflet and OpenStreetMap.")
    add_bullet("Implement Digital Proof-of-Service", "Develop a mobile-friendly driver manifest that validates checkpoint proximity via automated GPS geofencing, captures camera-based photo proof, and records immutable digital timestamps.")
    add_bullet("Standardize Exception & Skip Logging", "Provide structured mobile workflows allowing drivers to log valid operational obstructions (e.g., road repairs, impassable lanes) with photographic evidence and supervisory flagging.")
    add_bullet("Empower Citizens via Crowdsourced Hotspot Reporting", "Deploy an accessible public grievance portal where residents can report overflowing bins with device GPS pin-drops and track clearance progress via unique ticket IDs.")
    add_bullet("Generate Operational Performance Analytics", "Equip municipal administrators with comprehensive audit trails, ward-wise collection summaries, and route efficiency metrics to guide long-term urban policy.")

    doc.add_page_break()

    # ==========================================
    # CHAPTER 2: EXISTING SYSTEM
    # ==========================================
    add_heading_1("2. Existing System")
    
    add_heading_2("2.1 Overview of Current Manual Process")
    add_paragraph(
        "In the conventional municipal sanitation framework currently operational across most municipal wards, daily waste collection "
        "is managed through legacy manual methods. Every morning, drivers and ground sanitation crew assemble at the central zonal depot "
        "where supervisors verbally issue route instructions or hand over paper log sheets indicating the designated ward sector (e.g., "
        "Ward 1 Sangivalasa, Ward 2 Bheemili Coastal). The collection truck?typically a 2.5-ton tipper or 4-ton hydraulic compactor?departs "
        "the depot and follows historically established roadways."
    )
    add_paragraph(
        "Throughout the shift, drivers stop at public bin locations, where manual helpers shovel waste into the compactor. The driver notes "
        "completed stops on paper log sheets. If a road is impassable due to construction or traffic congestion, the driver simply bypasses "
        "the bin without any official recording. Upon shift completion, the vehicle dumps its load at the designated municipal landfill or "
        "transfer station, and the driver hands the handwritten log sheet back to the supervisor. Verification of whether each individual bin "
        "was actually emptied is entirely absent."
    )

    add_heading_2("2.2 Limitations of the Existing System")
    add_paragraph(
        "The manual, paper-based workflow suffers from critical structural vulnerabilities that compromise civic hygiene and administrative "
        "integrity:"
    )
    add_bullet("Total Lack of Real-Time Operational Visibility", "Supervisors stationed at municipal headquarters remain completely blind to vehicle whereabouts during the 8-hour shift. If a driver deviates from the planned route, idles excessively, or finishes early, no automated alert is generated.")
    add_bullet("Pervasive Unverified Service Claims", "Paper log sheets are inherently vulnerable to post-shift fabrication and falsification. Drivers can mark bins as 'cleared' even if they were bypassed, leading to uncollected garbage festering for days without management awareness.")
    add_bullet("Delayed Citizen Grievance Resolution Loop", "Residents experiencing overflowing bins must physically visit municipal ward offices or dial general helpline numbers. Complaints are recorded in physical registers, rarely correlated with active truck locations, and take days to reach field crew.")
    add_bullet("Inability to Identify Recurring Road Obstructions", "Because skipped stops are not digitally logged with categorized justifications, municipal engineers cannot identify recurring infrastructure issues?such as chronic illegal parking or prolonged pipeline trenching?that prevent waste clearance.")
    add_bullet("Absence of Empirical Data for Route Planning", "Without digital historical archives of pickup timestamps, route durations, and clearance volumes, urban planners cannot optimize collection schedules or balance vehicular workloads scientifically.")

    doc.add_page_break()

    # ==========================================
    # CHAPTER 3: PROPOSED SYSTEM
    # ==========================================
    add_heading_1("3. Proposed System")
    
    add_heading_2("3.1 Software-Only Architecture (Solution 2)")
    add_paragraph(
        "To decisively overcome the limitations of the existing manual system while avoiding the prohibitive financial and logistical "
        "pitfalls of IoT hardware sensors, this project implements Solution 2: The GPS-Based Fleet Tracking and Route Verification System "
        "(CivicClean GIS). This system operates on a pure software-only paradigm, utilizing standard web protocols, commodity smartphones, "
        "and client-side geolocation APIs."
    )
    add_paragraph(
        "Rather than installing dedicated microcontroller devices, cellular SIM modems, and ultrasonic sensors inside public bins, the "
        "system leverages the driver's smartphone as an intelligent edge sensor. The smartphone's integrated GPS receiver continuously "
        "transmits coordinates over cellular data networks to the municipal cloud server, while the built-in camera captures timestamped "
        "proof-of-service imagery upon collection. Concurrently, citizens use their own mobile devices to crowdsource hotspot reports, "
        "effectively replacing static hardware sensors with dynamic community monitoring."
    )

    add_heading_2("3.2 Key System Pillars")
    add_bullet("GIS Real-Time Fleet Tracking & Telemetry", "Sanitation supervisors monitor an interactive Leaflet/OpenStreetMap dashboard displaying live truck markers, color-coded route polylines, current travel speeds, and checkpoint status updates without requiring proprietary mapping API licenses.")
    add_bullet("Geofence-Enforced Proof-of-Service", "When a collection truck arrives at a scheduled bin checkpoint, the system computes the Haversine distance between the vehicle's live GPS coordinates and the pre-registered bin coordinates. Only when the driver is within an allowable proximity radius (e.g., <= 15 meters) does the app enable the 'Mark Collected' workflow, prompting a geotagged photo capture.")
    add_bullet("Structured Skip-Exception Management", "If a stop cannot be accessed due to physical barriers, the driver selects from standardized skip categories (e.g., road excavation, vehicle obstruction, hazardous waste) and inputs explanatory remarks. This immediately alerts supervisors and flags the checkpoint in red on the master GIS map.")
    add_bullet("Crowdsourced Citizen Grievance Portal", "Residents report waste accumulation by taking a photo, dropping a GPS pin on the municipal ward map, and selecting issue categories. The system automatically creates a trackable ticket (e.g., #CMP-2026-104) and dispatches nearby trucks.")
    add_bullet("Tamper-Resistant Audit Trail", "All collection events, skipped stops, and citizen tickets are immutably logged with precise timestamps, vehicle IDs, GPS accuracy tolerances, and verification hashes, providing audit-ready records for municipal administrative review.")

    add_heading_2("3.3 Comparative Advantages")
    add_paragraph(
        "The proposed software-only solution offers substantial qualitative and quantitative advantages over both legacy manual processes "
        "and IoT sensor networks:"
    )
    add_bullet("Zero Dedicated Hardware Capital Outlay", "Eliminates all expenditures associated with purchasing, installing, calibrating, and battery-servicing thousands of in-bin electronic sensors.")
    add_bullet("Zero Vandalism & Theft Risk", "Because all sensing resides securely on the driver's phone and citizen devices, municipal property in public thoroughfares is completely immune to physical destruction or theft.")
    add_bullet("Immediate Rapid Scalability", "Expanding the system to cover new wards, suburban colonies, or additional municipal vehicles requires merely entering route coordinates into the web portal; zero physical infrastructure deployment is needed.")
    add_bullet("Comprehensive Accountability & Transparency", "Establishes a closed digital feedback loop connecting all four civic stakeholders: residents, collection drivers, sanitation supervisors, and urban administrators.")

    doc.add_page_break()

    # ==========================================
    # CHAPTER 4: REQUIREMENTS SPECIFICATION
    # ==========================================
    add_heading_1("4. Requirements Specification")
    
    add_heading_2("4.1 Functional Requirements")
    add_paragraph(
        "Functional requirements define the core capabilities, operations, and services that the CivicClean system must provide to its "
        "authorized user roles:"
    )
    add_bullet("FR1: User Authentication & Role-Based Access Control", "The system shall authenticate municipal supervisors, field drivers, and citizens, granting tailored dashboard permissions based on role privileges.")
    add_bullet("FR2: Live GPS Telemetry Streaming", "The driver client shall continuously acquire device latitude/longitude coordinates and stream them to the server at regular intervals (every 3 to 10 seconds during active routes).")
    add_bullet("FR3: Interactive GIS Route & Fleet Display", "The supervisor dashboard shall render vehicle positions, route paths, checkpoint markers, and citizen complaint pins on an interactive OpenStreetMap canvas with dynamic popup metadata.")
    add_bullet("FR4: Geofence Validation & Proximity Matching", "The system shall calculate the real-time distance between the truck and the target checkpoint, validating whether the vehicle is within the designated collection geofence threshold (15m).")
    add_bullet("FR5: Digital Proof-of-Service Capture", "The driver portal shall capture collection timestamps, estimated waste weight, and geotagged photographic proof, transitioning the stop status to 'Collected' upon validation.")
    add_bullet("FR6: Structured Skip-Exception Logging", "The system shall enable drivers to bypass inaccessible stops by recording mandatory skip categories, driver remarks, and photo evidence, immediately alerting the supervisor.")
    add_bullet("FR7: Citizen Grievance Submission & Pin Placement", "The citizen portal shall allow residents to submit waste overflow complaints with contact details, ward selection, GPS pin placement, issue descriptions, and photo attachments.")
    add_bullet("FR8: Real-Time Grievance Status Tracking", "The system shall generate unique tracking ticket IDs and allow citizens to track complaint resolution progression (Submitted -> Dispatched -> Verified Cleared).")
    add_bullet("FR9: Proof-of-Service Audit Logging", "The backend shall maintain an immutable, chronologically ordered audit log of all completed collections, skipped stops, and supervisor overrides.")
    add_bullet("FR10: Operational Performance Reporting", "The system shall aggregate ward-wise waste weights, route completion percentages, and fuel efficiency metrics into downloadable reports.")

    add_heading_2("4.2 Non-Functional Requirements")
    add_paragraph(
        "Non-functional requirements specify quality attributes, behavioral constraints, and performance benchmarks that ensure system "
        "reliability, security, and responsiveness. These requirements are summarized in Table 4.1."
    )
    
    nfr_headers = ["Requirement ID", "Quality Dimension", "Specification & Benchmark"]
    nfr_rows = [
        ["NFR-01", "Performance & Latency", "API endpoints shall process and respond to telemetry and status requests within 250 milliseconds under standard network conditions. Map marker updates shall render within 100ms."],
        ["NFR-02", "Security & Data Protection", "All client-server communications shall use TLS/HTTPS encryption. Administrative endpoints shall enforce role-based access control (RBAC). Citizen phone numbers shall be masked."],
        ["NFR-03", "Availability & Uptime", "The cloud-hosted application server shall maintain 99.8% operational availability during active municipal collection hours (05:00 AM to 02:00 PM IST daily)."],
        ["NFR-04", "Usability & Accessibility", "The driver field portal shall feature a high-contrast, touch-optimized user interface operable with one-hand navigation. Text labels shall be clear and legible in direct outdoor sunlight."],
        ["NFR-05", "Geofence Precision", "The geofencing engine shall achieve spatial proximity accuracy within +/- 5 meters using smartphone Assisted GPS (A-GPS) and GLONASS satellite constellations."],
        ["NFR-06", "Scalability & Concurrency", "The backend architecture shall seamlessly scale to support up to 50 concurrent collection vehicles, 500 checkpoints, and 1,000 simultaneous citizen portal requests without memory saturation."],
        ["NFR-07", "Data Integrity & Auditability", "Audit logs and verification records shall be tamper-resistant, storing ISO-8601 timestamps and coordinate hashes that cannot be modified after initial submission."]
    ]
    add_table_custom("Table 4.1: Non-Functional Requirements Specification", nfr_headers, nfr_rows, [1.2, 1.8, 3.8])

    doc.add_page_break()

    # ==========================================
    # CHAPTER 5: HARDWARE AND SOFTWARE REQUIREMENTS
    # ==========================================
    add_heading_1("5. Hardware and Software Requirements")
    
    add_heading_2("5.1 Hardware Environment")
    add_paragraph(
        "A defining advantage of Solution 2 is its complete avoidance of specialized, expensive embedded hardware in bins. The system "
        "utilizes standard consumer hardware for development, field operation, and administrative supervision. Specifications are detailed "
        "in Table 5.1."
    )
    
    hw_headers = ["Role / Subsystem", "Hardware Component", "Minimum Specification", "Recommended Specification"]
    hw_rows = [
        ["Supervisor Workstation", "Processor", "Intel Core i3 / AMD Ryzen 3 (Quad Core)", "Intel Core i5 / AMD Ryzen 5 (Hexa Core)"],
        ["Supervisor Workstation", "System RAM", "4 GB DDR4", "8 GB DDR4 or higher"],
        ["Supervisor Workstation", "Display Monitor", "1366 x 768 Resolution", "1920 x 1080 Full HD (24-inch IPS)"],
        ["Supervisor Workstation", "Network Connectivity", "Broadband 10 Mbps", "High-speed Fiber Broadband 50+ Mbps"],
        ["Driver Field Device", "Handheld Smartphone", "Android 9.0+ / iOS 13+ Smartphone", "Android 12+ Smartphone with Octa-Core SoC"],
        ["Driver Field Device", "Sensors & Camera", "GPS/A-GPS receiver, 5 MP Camera", "Multi-GNSS (GPS+GLONASS), 12 MP Camera"],
        ["Driver Field Device", "Cellular Network", "3G / 4G LTE with Active SIM", "4G LTE / 5G High-Speed Data Plan"],
        ["Citizen Client Device", "Smartphone / PC", "Any smartphone, tablet, or desktop PC", "Modern smartphone with GPS location support"]
    ]
    add_table_custom("Table 5.1: Hardware Environment Specifications", hw_headers, hw_rows, [1.6, 1.4, 1.8, 2.0])

    add_heading_2("5.2 Software Environment")
    add_paragraph(
        "The software stack is designed around modern, lightweight, cross-platform technologies to ensure rapid deployment, zero licensing "
        "costs, and wide browser compatibility. Detailed specifications are given in Table 5.2."
    )
    
    sw_headers = ["Software Category", "Technology / Tool", "Version", "Purpose & Deployment Context"]
    sw_rows = [
        ["Operating System", "Microsoft Windows / Linux / macOS", "Windows 11 / Ubuntu 22.04 LTS", "Host OS for development and cloud server hosting"],
        ["Runtime Environment", "Node.js (LTS Engine)", "v24.13.0 / v20.x LTS", "High-performance asynchronous JavaScript server runtime"],
        ["Backend Framework", "HTTP / Express REST Framework", "v4.19+ / Built-in HTTP", "Lightweight REST API endpoints and static asset serving"],
        ["GIS Mapping Engine", "Leaflet.js & OpenStreetMap", "Leaflet v1.9.4", "Open-source, interactive tile rendering without API keys"],
        ["Frontend UI Stack", "HTML5, CSS3, JavaScript (ES6+)", "Modern Standards", "Responsive Single Page Application with tabbed role views"],
        ["Typography & Icons", "Google Inter, Font Awesome", "Font Awesome 6.4.0", "Modern UI vector glyphs and legible typography"],
        ["Web Browsers", "Google Chrome, Mozilla Firefox, Edge", "Chrome 120+, Edge 120+", "Cross-platform browser runtime for admin and clients"],
        ["Development Tools", "VS Code, Git, Python 3.14", "VS Code 1.95+", "Source code editing, version control, diagram scripts"]
    ]
    add_table_custom("Table 5.2: Software Environment Specifications", sw_headers, sw_rows, [1.4, 1.8, 1.3, 2.3])

    doc.add_page_break()
    print("Chapters 1 to 5 generated successfully.")
