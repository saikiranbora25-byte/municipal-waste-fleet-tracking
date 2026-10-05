# -*- coding: utf-8 -*-
"""Preliminary Pages: Title, Certificate, Declaration, Abstract, Lists"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def build_preliminary_pages(doc, helpers):
    add_paragraph = helpers['add_paragraph']
    
    # ==========================================
    # 1. TITLE / COVER PAGE (Exact Template Order & Font Sizes)
    # ==========================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(12)
    run_t = p_title.add_run("GPS-BASED FLEET TRACKING AND ROUTE VERIFICATION SYSTEM FOR MUNICIPAL SOLID WASTE MANAGEMENT")
    run_t.font.name = 'Times New Roman'
    run_t.font.size = Pt(14)               # Strict 14 pt Bold matching template
    run_t.bold = True
    run_t.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run("A Mini Project Report submitted as part of the academic requirements for the\n23CS4219- Software Engineering Laboratory\n\nBACHELOR OF TECHNOLOGY\nIN\nCOMPUTER SCIENCE AND ENGINEERING")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)               # Strict 12 pt matching template
    r_sub.bold = True
    
    # College Logo
    logo_path = 'report_assets/anits_logo.jpeg'
    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(8)
        p_logo.paragraph_format.space_after = Pt(8)
        p_logo.add_run().add_picture(logo_path, width=Inches(1.5))
        
    p_subm = doc.add_paragraph()
    p_subm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subm.paragraph_format.space_after = Pt(4)
    r_s1 = p_subm.add_run("Submitted by\n")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(12)
    r_s1.italic = True
    r_s2 = p_subm.add_run("SAIKIRAN BORA (A24126510006)\n\n")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(12)               # Strict 12 pt matching template
    r_s2.bold = True
    
    r_g1 = p_subm.add_run("Under the guidance of\n")
    r_g1.font.name = 'Times New Roman'
    r_g1.font.size = Pt(12)
    r_g1.italic = True
    r_g2 = p_subm.add_run("Prof. A. Rohini,\nDepartment of Computer Science and Engineering\n")
    r_g2.font.name = 'Times New Roman'
    r_g2.font.size = Pt(12)
    r_g2.bold = True
    
    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_dept.paragraph_format.space_before = Pt(12)
    p_dept.paragraph_format.space_after = Pt(0)
    r_dept = p_dept.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nANIL NEERUKONDA INSTITUTE OF TECHNOLOGY AND SCIENCES (UGC AUTONOMOUS)\n(Permanently Affiliated to AU, Approved by AICTE and Accredited by NBA & NAAC with 'A' Grade)\nSangivalasa, bheemili mandal, visakhapatnam dist.(A.P) 2026 - 2027")
    r_dept.font.name = 'Times New Roman'
    r_dept.font.size = Pt(12)              # Strict 12 pt matching template
    r_dept.bold = True
    
    doc.add_page_break()
    
    # ==========================================
    # 2. CERTIFICATE PAGE (Exact Template Text & Blank Left As-Is)
    # ==========================================
    p_c_head = doc.add_paragraph()
    p_c_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c_head.paragraph_format.space_after = Pt(2)
    r_ch = p_c_head.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nANIL NEERUKONDA INSTITUTE OF TECHNOLOGY AND SCIENCES\n(UGC AUTONOMOUS)\n(Affiliated to AU, Approved by AICTE and Accredited by NBA & NAAC with 'A' Grade)\nSangivalasa, bheemili mandal, visakhapatnam dist.(A.P)\n")
    r_ch.font.name = 'Times New Roman'
    r_ch.font.size = Pt(12)               # Strict 12 pt matching template
    r_ch.bold = True
    
    if os.path.exists(logo_path):
        p_c_logo = doc.add_paragraph()
        p_c_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c_logo.paragraph_format.space_before = Pt(4)
        p_c_logo.paragraph_format.space_after = Pt(6)
        p_c_logo.add_run().add_picture(logo_path, width=Inches(1.2))
        
    p_cert = doc.add_paragraph()
    p_cert.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert.paragraph_format.space_before = Pt(6)
    p_cert.paragraph_format.space_after = Pt(14)
    r_cert = p_cert.add_run("CERTIFICATE")
    r_cert.font.name = 'Times New Roman'
    r_cert.font.size = Pt(14)
    r_cert.bold = True
    
    # Exact template text with the blank left as-is:
    cert_text = (
        "This is to certify that the mini project entitled \"________________________\" has been successfully "
        "completed by the students of the Department of Computer Science and Engineering during the academic "
        "year 2026-2027 as part of the requirements for the 23CS4219-Software Engineering Laboratory."
    )
    p_c_body = doc.add_paragraph()
    p_c_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_c_body.paragraph_format.line_spacing = 1.3
    p_c_body.paragraph_format.space_after = Pt(30)
    r_cb = p_c_body.add_run(cert_text)
    r_cb.font.name = 'Times New Roman'
    r_cb.font.size = Pt(12)
    
    # Signature Table (Matching template layout)
    sig_table = doc.add_table(rows=2, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.rows[0].cells[0].width = Inches(3.2)
    sig_table.rows[0].cells[1].width = Inches(3.2)
    sig_table.rows[1].cells[0].width = Inches(3.2)
    sig_table.rows[1].cells[1].width = Inches(3.2)
    
    p_sig_top1 = sig_table.rows[0].cells[0].paragraphs[0]
    p_sig_top1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_sig_top1.paragraph_format.space_after = Pt(36)
    r_st1 = p_sig_top1.add_run("Faculty  Incharge")
    r_st1.font.name = 'Times New Roman'
    r_st1.font.size = Pt(12)
    r_st1.bold = True
    
    p_sig_top2 = sig_table.rows[0].cells[1].paragraphs[0]
    p_sig_top2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig_top2.paragraph_format.space_after = Pt(36)
    r_st2 = p_sig_top2.add_run("Head of the Department")
    r_st2.font.name = 'Times New Roman'
    r_st2.font.size = Pt(12)
    r_st2.bold = True
    
    p_sig1 = sig_table.rows[1].cells[0].paragraphs[0]
    p_sig1.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_s1 = p_sig1.add_run("Prof. A. Rohini")
    r_s1.font.name = 'Times New Roman'
    r_s1.font.size = Pt(12)
    r_s1.bold = False
    
    p_sig2 = sig_table.rows[1].cells[1].paragraphs[0]
    p_sig2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s2 = p_sig2.add_run("Prof. G. Srinivas")
    r_s2.font.name = 'Times New Roman'
    r_s2.font.size = Pt(12)
    r_s2.bold = False
    
    doc.add_page_break()
    
    # ==========================================
    # 3. DECLARATION PAGE
    # ==========================================
    p_d_head = doc.add_paragraph()
    p_d_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_d_head.paragraph_format.space_after = Pt(2)
    r_dh = p_d_head.add_run("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\nANIL NEERUKONDA INSTITUTE OF TECHNOLOGY AND SCIENCES\n(UGC AUTONOMOUS)\n(Affiliated to AU, Approved by AICTE and Accredited by NBA & NAAC with 'A' Grade)\nSangivalasa, bheemili mandal, visakhapatnam dist.(A.P)\n")
    r_dh.font.name = 'Times New Roman'
    r_dh.font.size = Pt(12)
    r_dh.bold = True
    
    p_decl = doc.add_paragraph()
    p_decl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_decl.paragraph_format.space_before = Pt(14)
    p_decl.paragraph_format.space_after = Pt(18)
    r_decl = p_decl.add_run("DECLARATION")
    r_decl.font.name = 'Times New Roman'
    r_decl.font.size = Pt(14)
    r_decl.bold = True
    
    decl_text = (
        "I, SAIKIRAN BORA (Roll No: A24126510006), student of III Year, I Semester B.Tech., Section - A, in the Department "
        "of Computer Science and Engineering from Anil Neerukonda Institute of Technology and Sciences (Autonomous), "
        "Visakhapatnam, hereby declare that the mini project work entitled \"GPS-BASED FLEET TRACKING AND ROUTE "
        "VERIFICATION SYSTEM FOR MUNICIPAL SOLID WASTE MANAGEMENT\" is an authentic record of independent work carried "
        "out by me under the guidance of Prof. A. Rohini, Department of Computer Science and Engineering.\n\n"
        "I further declare that this mini project report is submitted as part of the academic requirements for the "
        "23CS4219 - Software Engineering Laboratory during the academic year 2026-2027, and that the results embodied "
        "in this report have not been submitted to any other Institute or University for the award of any degree or diploma."
    )
    p_d_body = doc.add_paragraph()
    p_d_body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_d_body.paragraph_format.line_spacing = 1.3
    p_d_body.paragraph_format.space_after = Pt(36)
    r_db = p_d_body.add_run(decl_text)
    r_db.font.name = 'Times New Roman'
    r_db.font.size = Pt(12)
    
    p_d_sig = doc.add_paragraph()
    p_d_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ds = p_d_sig.add_run("SAIKIRAN BORA\nRoll No: A24126510006\nB.Tech. III Year, I Sem, Section - A\nDept. of CSE, ANITS")
    r_ds.font.name = 'Times New Roman'
    r_ds.font.size = Pt(12)
    r_ds.bold = True
    
    doc.add_page_break()
    # ==========================================
    # ==========================================
    # 4. ABSTRACT (150-250 WORDS)
    # ==========================================
    p_abs_h = doc.add_paragraph()
    p_abs_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_abs_h.paragraph_format.space_before = Pt(10)
    p_abs_h.paragraph_format.space_after = Pt(12)
    r_ah = p_abs_h.add_run("ABSTRACT")
    r_ah.font.name = 'Times New Roman'
    r_ah.font.size = Pt(14)
    r_ah.bold = True
    
    abstract_content = (
        "Municipal solid waste management in expanding urban sectors faces severe operational bottlenecks due to "
        "unmonitored paper-based collection schedules, unverified truck routes, fuel wastage, and delayed grievance "
        "redressal. While Internet-of-Things (IoT) hardware sensor systems have been proposed, their high procurement cost, "
        "frequent maintenance, battery exhaustion, and susceptibility to public theft/vandalism render them economically "
        "unviable for many municipal corporations. To overcome these critical barriers, this project implements a pure "
        "software-only solution: the GPS-Based Fleet Tracking and Route Verification System (CivicClean GIS).\n\n"
        "The system completely eliminates custom in-bin hardware by utilizing commodity driver smartphones and modern web "
        "browsers. It delivers four core architectural modules: (1) an interactive real-time GIS fleet tracking map for sanitation "
        "supervisors powered by Leaflet and OpenStreetMap; (2) a mobile-responsive driver manifest featuring automated GPS "
        "geofence proximity matching, camera-based proof-of-service photo uploads, and structured skip-exception logging; "
        "(3) a public crowdsourced citizen reporting portal allowing residents to submit geotagged photos of overflowing bins "
        "and track grievance tickets in real time; and (4) a tamper-resistant proof-of-service audit trail that cross-references "
        "collection timestamps and location tolerances. Developed using a robust Node.js/Express REST backend and a responsive "
        "HTML5/CSS3/JavaScript frontend, the system was thoroughly evaluated across municipal test wards in Visakhapatnam. The "
        "validation results confirmed 100% route verifiability, sub-10-meter geofence precision, instant exception capture, and "
        "zero dedicated hardware expense, offering a scalable and accountable solution for modern urban sanitation governance."
    )
    p_abs_b = doc.add_paragraph()
    p_abs_b.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs_b.paragraph_format.line_spacing = 1.2
    p_abs_b.paragraph_format.space_after = Pt(12)
    r_ab = p_abs_b.add_run(abstract_content)
    r_ab.font.name = 'Times New Roman'
    r_ab.font.size = Pt(12)
    
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_k1 = p_kw.add_run("Keywords: ")
    r_k1.font.name = 'Times New Roman'
    r_k1.font.size = Pt(11)
    r_k1.bold = True
    r_k2 = p_kw.add_run("Municipal Solid Waste Management, GPS Fleet Tracking, Route Verification, Proof-of-Service, Geofencing, Citizen Grievance Redressal, Zero-Hardware Architecture.")
    r_k2.font.name = 'Times New Roman'
    r_k2.font.size = Pt(11)
    r_k2.italic = True
    
    doc.add_page_break()
    
    # ==========================================
    # 5. TABLE OF CONTENTS
    # ==========================================
    p_toc_h = doc.add_paragraph()
    p_toc_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_h.paragraph_format.space_before = Pt(8)
    p_toc_h.paragraph_format.space_after = Pt(10)
    r_th = p_toc_h.add_run("TABLE OF CONTENTS")
    r_th.font.name = 'Times New Roman'
    r_th.font.size = Pt(14)
    r_th.bold = True
    
    toc_items = [
        ("Certificate", "ii"),
        ("Declaration", "iii"),
        ("Abstract", "iv"),
        ("List of Figures", "vi"),
        ("List of Tables", "vii"),
        ("1. Introduction", "1"),
        ("   1.1 Background", "1"),
        ("   1.2 Problem Statement", "2"),
        ("   1.3 Motivation", "2"),
        ("   1.4 Objectives", "3"),
        ("2. Existing System", "4"),
        ("   2.1 Overview of Current Manual Process", "4"),
        ("   2.2 Limitations of the Existing System", "5"),
        ("3. Proposed System", "6"),
        ("   3.1 Software-Only Architecture (Solution 2)", "6"),
        ("   3.2 Key System Pillars", "7"),
        ("   3.3 Comparative Advantages", "8"),
        ("4. Requirements Specification", "9"),
        ("   4.1 Functional Requirements", "9"),
        ("   4.2 Non-Functional Requirements", "10"),
        ("5. Hardware and Software Requirements", "12"),
        ("   5.1 Hardware Environment", "12"),
        ("   5.2 Software Environment", "12"),
        ("6. Data Flow Diagram", "14"),
        ("   6.1 Level 0 Context DFD", "14"),
        ("   6.2 Level 1 DFD", "15"),
        ("   6.3 Level 2 DFD (Proof-of-Service Decomposition)", "17"),
        ("7. System Architecture", "19"),
        ("   7.1 3-Tier Client-Server Architecture", "19"),
        ("   7.2 Architectural Layer Analysis", "20"),
        ("8. Module Description", "22"),
        ("   8.1 Supervisor Command Hub & GIS Fleet Tracking", "22"),
        ("   8.2 Driver Mobile Field Manifest & Geofencing", "23"),
        ("   8.3 Proof-of-Service Verification & Tamper-Resistant Audit", "24"),
        ("   8.4 Citizen Grievance Redressal & Hotspot Triage", "25"),
        ("9. System Analysis", "26"),
        ("   9.1 Use Case Diagram", "26"),
        ("   9.2 Use Case Specifications", "27"),
        ("10. UML Design", "30"),
        ("   10.1 Class Diagram", "30"),
        ("   10.2 Sequence Diagram", "32"),
        ("   10.3 Activity Diagram", "34"),
        ("   10.4 State Diagram", "36"),
        ("   10.5 Component Diagram", "38"),
        ("   10.6 Deployment Diagram", "40"),
        ("11. Testing", "42"),
        ("   11.1 Test Plan & Methodology", "42"),
        ("   11.2 System Test Cases Execution Matrix", "43"),
        ("12. Results / Screenshots", "46"),
        ("   12.1 Home Page & Portal Switcher", "46"),
        ("   12.2 Role-Based Authentication Interface", "47"),
        ("   12.3 Supervisor Live GIS Fleet Map", "48"),
        ("   12.4 Driver Mobile Field Manifest & Proof Upload", "49"),
        ("   12.5 Citizen Grievance Portal & Live Ticket Tracker", "50"),
        ("   12.6 Proof-of-Service Audit Trail", "51"),
        ("   12.7 Fleet Waste Clearance Analytics Report", "52"),
        ("13. Limitations", "53"),
        ("14. Conclusion", "54"),
        ("15. References", "55"),
        ("Appendix", "56"),
        ("   Appendix A: System User Manual", "56"),
        ("   Appendix B: REST API Payloads & Data Schemas", "57"),
        ("   Appendix C: Project Book Spiral Binding Guidelines", "58")
    ]
    
    toc_table = doc.add_table(rows=len(toc_items), cols=2)
    toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (title, page) in enumerate(toc_items):
        r = toc_table.rows[idx]
        r.cells[0].width = Inches(5.6)
        r.cells[1].width = Inches(0.8)
        
        p0 = r.cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.05
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(title)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        if not title.startswith("   "):
            r0.bold = True
            
        p1 = r.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(page)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        if not title.startswith("   "):
            r1.bold = True
            
    doc.add_page_break()
    
    # ==========================================
    # 6. LIST OF FIGURES
    # ==========================================
    p_lof_h = doc.add_paragraph()
    p_lof_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lof_h.paragraph_format.space_before = Pt(8)
    p_lof_h.paragraph_format.space_after = Pt(10)
    r_lfh = p_lof_h.add_run("LIST OF FIGURES")
    r_lfh.font.name = 'Times New Roman'
    r_lfh.font.size = Pt(14)
    r_lfh.bold = True
    
    figures = [
        ("Figure 6.1", "Level 0 Context Data Flow Diagram (DFD)", "14"),
        ("Figure 6.2", "Level 1 Data Flow Diagram (DFD)", "16"),
        ("Figure 6.3", "Level 2 DFD (Process 3.0 Proof-of-Service Decomposition)", "18"),
        ("Figure 7.1", "3-Tier System Architecture Diagram", "20"),
        ("Figure 9.1", "System Use Case Diagram", "26"),
        ("Figure 10.1", "UML Class Diagram", "31"),
        ("Figure 10.2", "UML Sequence Diagram (Proof-of-Service Verification)", "33"),
        ("Figure 10.3", "UML Activity Diagram (Driver Daily Collection Workflow)", "35"),
        ("Figure 10.4", "UML State Diagram (Stop Checkpoint Lifecycle)", "37"),
        ("Figure 10.5", "UML Component Diagram", "39"),
        ("Figure 10.6", "UML Deployment Diagram", "41"),
        ("Figure 12.1", "Application Home Page & Portal Switcher View", "46"),
        ("Figure 12.2", "Role-Based User Authentication Interface", "47"),
        ("Figure 12.3", "Supervisor Live GIS Fleet Tracking Dashboard", "48"),
        ("Figure 12.4", "Driver Mobile Field Manifest & Proof Upload Interface", "49"),
        ("Figure 12.5", "Citizen Grievance Submission Portal & Live Ticket Tracker", "50"),
        ("Figure 12.6", "Proof-of-Service Tamper-Resistant Audit Trail", "51"),
        ("Figure 12.7", "Fleet Waste Clearance Analytics & Performance Report", "52")
    ]
    
    lof_table = doc.add_table(rows=len(figures), cols=3)
    lof_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (f_num, f_title, f_page) in enumerate(figures):
        r = lof_table.rows[idx]
        r.cells[0].width = Inches(1.3)
        r.cells[1].width = Inches(4.3)
        r.cells[2].width = Inches(0.8)
        
        p0 = r.cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.05
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(f_num)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True
        
        p1 = r.cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(f_title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        
        p2 = r.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(f_page)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.bold = True
        
    doc.add_page_break()
    
    # ==========================================
    # 7. LIST OF TABLES
    # ==========================================
    p_lot_h = doc.add_paragraph()
    p_lot_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_lot_h.paragraph_format.space_before = Pt(8)
    p_lot_h.paragraph_format.space_after = Pt(10)
    r_lth = p_lot_h.add_run("LIST OF TABLES")
    r_lth.font.name = 'Times New Roman'
    r_lth.font.size = Pt(14)
    r_lth.bold = True
    
    tables = [
        ("Table 4.1", "Non-Functional Requirements Specification", "11"),
        ("Table 5.1", "Hardware Environment Specifications", "13"),
        ("Table 5.2", "Software Environment Specifications", "13"),
        ("Table 9.1", "Use Case Specification: UC1 View Assigned Route Manifest", "27"),
        ("Table 9.2", "Use Case Specification: UC3 Mark Stop Collected with Proof", "28"),
        ("Table 9.3", "Use Case Specification: UC5 Report Waste Hotspot", "29"),
        ("Table 11.1", "System Test Cases Matrix & Validation Results", "44")
    ]
    
    lot_table = doc.add_table(rows=len(tables), cols=3)
    lot_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (t_num, t_title, t_page) in enumerate(tables):
        r = lot_table.rows[idx]
        r.cells[0].width = Inches(1.3)
        r.cells[1].width = Inches(4.3)
        r.cells[2].width = Inches(0.8)
        
        p0 = r.cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.05
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(t_num)
        r0.font.name = 'Times New Roman'
        r0.font.size = Pt(11)
        r0.bold = True
        
        p1 = r.cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(t_title)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(11)
        
        p2 = r.cells[2].paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(t_page)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(11)
        r2.bold = True
        
    doc.add_page_break()
    print("Preliminary pages generated successfully.")
