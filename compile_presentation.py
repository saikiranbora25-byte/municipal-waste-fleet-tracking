# -*- coding: utf-8 -*-
"""
Main Compiler for CivicClean Project Presentation (16 Slides)
Generates: CivicClean_Municipal_Waste_Fleet_Tracking_Presentation.pptx
Author: Saikiran Bora (Roll No: A24126510006)
Branch: Computer Science & Engineering | College: ANITS
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

print("Initializing Presentation Compilation with 12pt Typography...")

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

C_DARK_BG   = RGBColor(11, 19, 43)      # #0B132B
C_LIGHT_BG  = RGBColor(248, 250, 252)   # #F8FAFC
C_WHITE     = RGBColor(255, 255, 255)
C_PRIMARY   = RGBColor(37, 99, 235)     # #2563EB
C_PRIMARY_DK= RGBColor(30, 58, 138)     # #1E3A8A
C_SUCCESS   = RGBColor(16, 185, 129)    # #10B981
C_DANGER    = RGBColor(239, 68, 68)     # #EF4444
C_WARNING   = RGBColor(245, 158, 11)    # #F59E0B
C_TEXT_MAIN = RGBColor(15, 23, 42)      # #0F172A
C_TEXT_SUB  = RGBColor(71, 85, 105)     # #475569
C_TEXT_MUTED= RGBColor(100, 116, 139)   # #64748B
C_BORDER    = RGBColor(226, 232, 240)   # #E2E8F0
C_CARD_FILL = RGBColor(255, 255, 255)

def apply_background(slide, color=C_LIGHT_BG):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg

def add_header(slide, title_text, category="MUNICIPAL SOLID WASTE MANAGEMENT | SOLUTION 2", slide_num=None):
    # Category Pill
    cat_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.38), Inches(4.8), Inches(0.34))
    cat_box.fill.solid()
    cat_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
    cat_box.line.color.rgb = RGBColor(191, 219, 254)
    cat_box.line.width = Pt(1)
    tf_c = cat_box.text_frame
    tf_c.vertical_anchor = MSO_ANCHOR.MIDDLE
    p_c = tf_c.paragraphs[0]
    p_c.text = category.upper()
    p_c.font.name = "Segoe UI"
    p_c.font.size = Pt(9.5)
    p_c.font.bold = True
    p_c.font.color.rgb = C_PRIMARY
    
    # Title Text
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.68))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_MAIN

    # Divider
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.44), Inches(11.733), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = C_BORDER
    div.line.fill.background()

    # Footer
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf_f = f_box.text_frame
    p_f = tf_f.paragraphs[0]
    p_f.font.name = "Segoe UI"
    p_f.font.size = Pt(9)
    p_f.font.color.rgb = C_TEXT_MUTED
    p_f.text = "Saikiran Bora (Roll No: A24126510006) | Branch: CSE, ANITS  •  CivicClean GIS: GPS Fleet Tracking & Route Verification Platform"
    if slide_num:
        p_f.text += f"  •  Slide {slide_num} of 16"

def add_card(slide, left, top, width, height, title="", border_color=C_BORDER, fill_color=C_CARD_FILL, title_color=C_PRIMARY_DK):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1.2)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Segoe UI"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.space_after = Pt(6)
    return card, tf

helpers = {
    'prs': prs,
    'blank_layout': blank_layout,
    'apply_background': apply_background,
    'add_header': add_header,
    'add_card': add_card,
    'C_DARK_BG': C_DARK_BG,
    'C_LIGHT_BG': C_LIGHT_BG,
    'C_WHITE': C_WHITE,
    'C_PRIMARY': C_PRIMARY,
    'C_PRIMARY_DK': C_PRIMARY_DK,
    'C_SUCCESS': C_SUCCESS,
    'C_DANGER': C_DANGER,
    'C_WARNING': C_WARNING,
    'C_TEXT_MAIN': C_TEXT_MAIN,
    'C_TEXT_SUB': C_TEXT_SUB,
    'C_TEXT_MUTED': C_TEXT_MUTED,
    'C_BORDER': C_BORDER,
    'C_CARD_FILL': C_CARD_FILL
}

import ppt_part1 as pptt_part1
import ppt_part2 as pptt_part2
import ppt_part3 as pptt_part3
import ppt_part4 as pptt_part4

print("Compiling Slides 1 to 4...")
pptt_part1.build_slides_1_to_4(prs, helpers)

print("Compiling Slides 5 to 8...")
pptt_part2.build_slides_5_to_8(prs, helpers)

print("Compiling Slides 9 to 12...")
pptt_part3.build_slides_9_to_12(prs, helpers)

print("Compiling Slides 13 to 16...")
pptt_part4.build_slides_13_to_16(prs, helpers)

output_ppt = "CivicClean_Municipal_Waste_Fleet_Tracking_Presentation.pptx"
prs.save(output_ppt)
print(f"Presentation successfully compiled and saved to: {output_ppt}")
print(f"Total slides compiled: {len(prs.slides)}")
